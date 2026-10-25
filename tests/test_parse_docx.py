"""DOCX parser. Fixtures are built in the test."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from ragcli.errors import EXIT_DOCUMENT, RagError
from ragcli.parse import parse_file
from tests.parse_support import assert_text_chunks, bomb_png, solid_image


def test_paragraphs_table_and_image_keep_body_order(tmp_path: Path) -> None:
    png_path = tmp_path / "shot.png"
    png_path.write_bytes(solid_image(140, 90, (12, 80, 160)))
    path = tmp_path / "note.docx"
    document = Document()
    document.add_heading("營業額", level=1)
    document.add_paragraph("前文在圖的前面。")
    document.add_picture(str(png_path), width=Inches(1.2))
    document.add_paragraph("後文在圖的後面。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "名稱"
    table.cell(0, 1).text = "金額"
    table.cell(1, 0).text = "茶"
    table.cell(1, 1).text = "12"
    document.save(path)

    result = parse_file(path)
    assert result.document.media_type == "docx"
    assert result.document.parser == "python-docx"
    assert result.document.page_count is None
    assert result.document.title is None
    assert_text_chunks(result)
    texts = [chunk.text for chunk in result.chunks]
    joined = "\n".join(texts)
    assert "營業額" in joined
    assert "前文在圖的前面。" in joined
    assert "後文在圖的後面。" in joined
    assert "名稱\t金額" in joined
    assert "茶\t12" in joined
    assert all(chunk.section == "營業額" for chunk in result.chunks)
    assert len(result.images) == 1
    image = result.images[0]
    assert image.kind == "embedded"
    assert image.status == "pending_ocr"
    assert image.page is None
    assert image.png_bytes is not None
    before = next(chunk for chunk in result.chunks if "前文" in chunk.text)
    after = next(chunk for chunk in result.chunks if "後文" in chunk.text)
    assert before.ordinal < image.preceding_chunk_count
    assert after.ordinal >= image.preceding_chunk_count
    assert Image.open(io.BytesIO(image.png_bytes)).size == (140, 90)
    assert list(tmp_path.iterdir()) == [png_path, path] or set(tmp_path.iterdir()) == {png_path, path}


def test_bomb_image_is_skipped_and_text_remains(tmp_path: Path, monkeypatch) -> None:
    png_path = tmp_path / "ok.png"
    png_path.write_bytes(solid_image(30, 20, (1, 2, 3)))
    path = tmp_path / "bomb.docx"
    document = Document()
    document.add_paragraph("圖旁邊仍有文字。")
    document.add_picture(str(png_path))
    document.save(path)
    _replace_media(path, bomb_png(8000, 5000))
    monkeypatch.setattr(
        "ragcli.parse.pixels.Image.open",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("decoded")),
    )

    result = parse_file(path)
    assert any("圖旁邊仍有文字。" in chunk.text for chunk in result.chunks)
    assert len(result.images) == 1
    image = result.images[0]
    assert image.route == "skipped_bomb"
    assert image.status == "skipped_bomb"
    assert image.png_bytes is None
    assert image.width * image.height > 20_000_000


def test_corrupt_docx_is_exit_5(tmp_path: Path) -> None:
    path = tmp_path / "bad.docx"
    path.write_bytes(b"not a docx")
    with pytest.raises(RagError) as exc:
        parse_file(path)
    assert exc.value.code == EXIT_DOCUMENT


def _replace_media(path: Path, payload: bytes) -> None:
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            data = payload if info.filename.startswith("word/media/") else source.read(info)
            target.writestr(info, data)
    path.write_bytes(buffer.getvalue())

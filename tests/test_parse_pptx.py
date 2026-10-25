"""PPTX parser. Fixtures are built in the test."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from ragcli.errors import EXIT_DOCUMENT, RagError
from ragcli.parse import ParseLimits, parse_file
from tests.parse_support import assert_text_chunks, bomb_png, solid_image


def _blank(presentation: Presentation):
    index = 6 if len(presentation.slide_layouts) > 6 else len(presentation.slide_layouts) - 1
    return presentation.slide_layouts[index]


def test_slide_text_picture_and_notes(tmp_path: Path) -> None:
    png_path = tmp_path / "shot.png"
    png_path.write_bytes(solid_image(160, 100, (30, 90, 40)))
    path = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(_blank(presentation))
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(6), Inches(1))
    box.text_frame.text = "投影片講營業額同 revenue。"
    slide.shapes.add_picture(str(png_path), Inches(0.5), Inches(1.6), Inches(2))
    table_shape = slide.shapes.add_table(2, 2, Inches(0.5), Inches(3.4), Inches(4), Inches(1))
    table_shape.table.cell(0, 0).text = "名稱"
    table_shape.table.cell(0, 1).text = "金額"
    table_shape.table.cell(1, 0).text = "茶"
    table_shape.table.cell(1, 1).text = "12"
    slide.notes_slide.notes_text_frame.text = "備註只在 notes。"
    presentation.save(path)

    result = parse_file(path)
    assert result.document.media_type == "pptx"
    assert result.document.parser == "python-pptx"
    assert result.document.page_count == 1
    assert_text_chunks(result)
    body = [chunk for chunk in result.chunks if chunk.section != "notes"]
    notes = [chunk for chunk in result.chunks if chunk.section == "notes"]
    assert body and notes
    joined = "\n".join(chunk.text for chunk in body)
    assert "投影片講營業額同 revenue。" in joined
    assert "名稱\t金額" in joined
    assert "茶\t12" in joined
    assert notes[0].text == "備註只在 notes。"
    assert notes[0].page_start == 1
    assert len(result.images) == 1
    image = result.images[0]
    assert image.kind == "embedded"
    assert image.page == 1
    assert image.status == "pending_ocr"
    assert image.png_bytes is not None
    assert Image.open(io.BytesIO(image.png_bytes)).size == (160, 100)
    assert image.preceding_chunk_count <= notes[0].ordinal


def test_bomb_picture_is_not_decoded(tmp_path: Path, monkeypatch) -> None:
    png_path = tmp_path / "ok.png"
    png_path.write_bytes(solid_image(20, 20, (4, 5, 6)))
    path = tmp_path / "bomb.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(_blank(presentation))
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(4), Inches(0.5))
    box.text_frame.text = "還有文字。"
    slide.shapes.add_picture(str(png_path), Inches(0.5), Inches(1.2))
    presentation.save(path)
    _replace_media(path, bomb_png(9000, 4000))
    monkeypatch.setattr(
        "ragcli.parse.pixels.Image.open",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("decoded")),
    )

    result = parse_file(path)
    assert any("還有文字。" in chunk.text for chunk in result.chunks)
    assert result.images[0].route == "skipped_bomb"
    assert result.images[0].png_bytes is None


def test_slide_limit_is_exit_5(tmp_path: Path) -> None:
    path = tmp_path / "two.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(_blank(presentation))
    presentation.slides.add_slide(_blank(presentation))
    presentation.save(path)
    with pytest.raises(RagError) as exc:
        parse_file(path, ParseLimits(max_pages=1))
    assert exc.value.code == EXIT_DOCUMENT


def _replace_media(path: Path, payload: bytes) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            data = payload if info.filename.startswith("ppt/media/") else source.read(info)
            target.writestr(info, data)
    path.write_bytes(buffer.getvalue())

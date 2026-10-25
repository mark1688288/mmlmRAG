"""Ingest pipeline against fake Ollama and fake OCR.

SQLite and Chroma are real, under tmp_path. No socket, no pull, and no vm_stat.
"""

from __future__ import annotations

import hashlib
import io
import json
import logging
import sqlite3
import stat
import zipfile
from pathlib import Path

import httpx
import pytest
from docx import Document as DocxDocument
from docx.shared import Inches

from ragcli.config import Config, load_config
from ragcli.errors import (
    EXIT_DISK,
    EXIT_DOCUMENT,
    EXIT_MODEL_MISSING,
    EXIT_NOT_FOUND,
    EXIT_NO_CHUNKS,
    EXIT_OK,
    EXIT_OLLAMA_UNREACHABLE,
    EXIT_RESOURCES,
    EXIT_USAGE,
    RagError,
)
from ragcli.ingest import run_ingest
from ragcli.ollama import NINE_B_MODEL
from ragcli.parse import parse_file
from ragcli.parse.model import (
    PARSER_VERSION,
    ParsedDocument,
    ParsedImage,
    ParseResult,
    image_id_for,
    storage_relpath,
)
from ragcli.store.chroma import ChromaIndex, IndexRecord, collection_name
from ragcli.store.sqlite import (
    ActiveIndex,
    Chunk,
    Document,
    ImageRecord,
    SqliteStore,
    chunk_id_for,
)
from tests.parse_support import bomb_png, solid_image
from tests.test_caption import FakeOllama

EMBED = "qwen3-embedding:0.6b"
FOUR_B = "qwen3.5:4b"
AVAILABLE = 20_000_000_000
WHEN = "2026-10-26T00:00:00+00:00"
VECTOR = (1.0, 0.0, 0.0, 4.0)
NATIVE = 4
WIDTH = 4
DOC_A = "a" * 64
DOC_B = "b" * 64
DOC_C = "c" * 64


class IngestFake(FakeOllama):
    """Fake that answers version and returns native-width embeddings."""

    def __init__(self) -> None:
        super().__init__()
        self.tag_names = [EMBED, NINE_B_MODEL, FOUR_B]
        self.vector = list(VECTOR)
        self.embed_inputs: list[list[str]] = []
        self.loaded_during_embed: list[list[str]] = []
        self.down = False

    def handler(self, request: httpx.Request) -> httpx.Response:
        if request.url.path == "/api/version":
            self.seen.append((request.method, request.url.path, None))
            if self.down:
                raise httpx.ConnectError("down")
            return httpx.Response(200, json={"version": "0.0.0"})
        return super().handler(request)

    def _embed(self, body: dict) -> httpx.Response:
        if body.get("input") == ".":
            self._drop(body["model"])
            return httpx.Response(200, json={"embeddings": [[0.0]]})
        self.loaded_during_embed.append([str(item.get("name")) for item in self.loaded])
        texts = list(body["input"])
        self.embed_inputs.append(texts)
        self._show(body["model"], body["options"]["num_ctx"])
        return httpx.Response(
            200, json={"embeddings": [list(self.vector) for _text in texts]}
        )

    def document_inputs(self) -> list[str]:
        return [text for batch in self.embed_inputs for text in batch]

    def embed_bodies(self) -> list[dict]:
        return [
            body
            for body in self.bodies("/api/embed")
            if isinstance(body, dict) and isinstance(body.get("input"), list)
        ]


def write_rag(
    directory: Path,
    *,
    batch_size: int = 8,
    dimensions: int = WIDTH,
    native: int = NATIVE,
    min_free: int = 1,
) -> Config:
    path = directory / "rag.toml"
    path.write_text(
        "\n".join(
            [
                "[paths]",
                'data_dir = "data"',
                "[embed]",
                f'model = "{EMBED}"',
                f"native_dimensions = {native}",
                f"dimensions = {dimensions}",
                "num_ctx = 2048",
                f"batch_size = {batch_size}",
                "[limits]",
                f"min_free_bytes = {min_free}",
                "",
            ]
        ),
        encoding="utf-8",
    )
    return load_config(path)


def refuse_ocr(*_args, **_kwargs):
    raise AssertionError("不應呼叫 OCR")


def ocr_engine(text: str, confidence: float = 0.8):
    def engine(_image, *, languages, recognition_level, confidence_threshold):
        del languages, recognition_level, confidence_threshold
        if text == "":
            return []
        return [(text, confidence, (0.1, 0.8, 0.2, 0.1))]

    return engine


def run(config: Config, fake: IngestFake | None, path: Path | None, **kwargs) -> int:
    client = None if fake is None else fake.client()
    return run_ingest(
        config,
        path=path,
        client=client,
        available=kwargs.pop("available", AVAILABLE),
        ocr_engine=kwargs.pop("ocr_engine", refuse_ocr),
        clock=kwargs.pop("clock", lambda: WHEN),
        **kwargs,
    )


def open_store(config: Config) -> SqliteStore:
    return SqliteStore.open(config.data_dir)


def write_png(path: Path, width: int = 140, height: int = 90) -> None:
    path.write_bytes(solid_image(width, height, (12, 80, 160)))


def write_docx(path: Path, png: Path) -> None:
    document = DocxDocument()
    document.add_heading("營業額", level=1)
    document.add_paragraph("前文在圖的前面。")
    document.add_picture(str(png), width=Inches(1.2))
    document.add_paragraph("後文在圖的後面。")
    document.save(path)


def replace_media(path: Path, payload: bytes) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            data = payload if info.filename.startswith("word/media/") else source.read(info)
            target.writestr(info, data)
    path.write_bytes(buffer.getvalue())


def seed_document(
    store: SqliteStore,
    doc_id: str,
    source: str,
    text: str,
    *,
    parser_version: str = PARSER_VERSION,
    embed_state: str = "pending",
) -> None:
    store.insert_document(
        Document(
            doc_id=doc_id,
            source_path=source,
            filename=Path(source).name,
            media_type="text",
            byte_size=8,
            sha256=doc_id,
            parser="text",
            parser_version=parser_version,
            ingested_at=WHEN,
            status="parsed",
        )
    )
    store.insert_chunk(
        Chunk(
            chunk_id=chunk_id_for(doc_id, 0),
            doc_id=doc_id,
            ordinal=0,
            text=text,
            kind="text",
            token_estimate=4,
            embed_state=embed_state,
        )
    )


def count_parse(monkeypatch) -> dict[str, int]:
    calls = {"n": 0}
    real = parse_file

    def spy(path, limits=None):
        calls["n"] += 1
        return real(path, limits)

    monkeypatch.setattr("ragcli.parse.parse_file", spy)
    return calls


def test_docx_runs_seven_stages_in_order(tmp_path, monkeypatch, capsys):
    config = write_rag(tmp_path)
    png = tmp_path / "shot.png"
    write_png(png)
    document = tmp_path / "note.docx"
    write_docx(document, png)
    fake = IngestFake()
    fake.replies = [{"role": "assistant", "content": "這是一張圖表。", "thinking": "不要採用"}]
    events: list[str] = []
    seen_at_ocr: dict = {}
    seen_at_chat: dict = {}
    real = parse_file

    def parse_spy(path, limits=None):
        events.append("parse")
        return real(path, limits)

    monkeypatch.setattr("ragcli.parse.parse_file", parse_spy)

    def engine(_image, *, languages, recognition_level, confidence_threshold):
        del languages, recognition_level, confidence_threshold
        events.append("ocr")
        store = open_store(config)
        try:
            saved = store.list_documents()[0]
            chunks = store.chunks_for_document(saved.doc_id)
            image_row = store.images_for_document(saved.doc_id)[0]
            seen_at_ocr["status"] = saved.status
            seen_at_ocr["kinds"] = [chunk.kind for chunk in chunks]
            seen_at_ocr["image_status"] = image_row.status
            image_path = store.image_path(saved.doc_id, image_row.image_id)
            seen_at_ocr["png"] = image_path.is_file()
            seen_at_ocr["mode"] = stat.S_IMODE(image_path.stat().st_mode)
        finally:
            store.close()
        return [("圖表說明", 0.8, (0.1, 0.8, 0.2, 0.1))]

    original = fake.handler

    def spy(request: httpx.Request) -> httpx.Response:
        path = request.url.path
        body = json_body(request)
        if path == "/api/version":
            events.append("version")
        elif path == "/api/tags":
            events.append("tags")
        elif path == "/api/chat":
            events.append("chat")
            store = open_store(config)
            try:
                saved = store.list_documents()[0]
                figure = next(
                    chunk
                    for chunk in store.chunks_for_document(saved.doc_id)
                    if chunk.kind == "figure"
                )
                anchor = store.get_chunk(figure.anchor_chunk_id or "")
                seen_at_chat["text"] = figure.text
                seen_at_chat["state"] = figure.embed_state
                seen_at_chat["anchor_text"] = "" if anchor is None else anchor.text
            finally:
                store.close()
        elif path == "/api/embed" and isinstance(body, dict) and isinstance(body.get("input"), list):
            events.append("embed")
        return original(request)

    fake.handler = spy  # type: ignore[method-assign]
    logger = logging.getLogger("ragcli")
    previous = logger.level
    logger.setLevel(logging.INFO)
    try:
        code = run(config, fake, document, ocr_engine=engine)
    finally:
        logger.setLevel(previous)

    captured = capsys.readouterr()
    assert code == EXIT_OK
    assert events.index("version") < events.index("tags") < events.index("parse")
    assert events.index("parse") < events.index("ocr") < events.index("chat")
    assert events.index("chat") < events.index("embed")
    assert seen_at_ocr["status"] == "parsed"
    assert set(seen_at_ocr["kinds"]) == {"text"}
    assert seen_at_ocr["image_status"] == "pending_ocr"
    assert seen_at_ocr["png"] is True
    assert seen_at_ocr["mode"] == 0o600
    assert seen_at_chat["state"] == "blocked"
    assert "描述" not in seen_at_chat["text"]
    assert "前文在圖的前面" in seen_at_chat["text"]
    assert "後文在圖的後面" in seen_at_chat["text"]
    assert "圖中文字：圖表說明" in seen_at_chat["text"]
    assert "前文" in seen_at_chat["anchor_text"]

    chat = fake.bodies("/api/chat")[0]
    assert chat["think"] is False
    assert chat["keep_alive"] == "30m"
    assert chat["options"]["num_ctx"] == 4096
    assert chat["options"]["num_predict"] == 400
    assert chat["options"]["temperature"] == 0.1
    embed_body = fake.embed_bodies()[0]
    assert embed_body["truncate"] is False
    assert "dimensions" not in embed_body
    assert embed_body["keep_alive"] == "30m"
    assert embed_body["options"]["num_ctx"] == 2048
    assert embed_body["model"] == EMBED
    assert all(not text.startswith("Instruct:") for text in fake.document_inputs())
    assert any("描述：這是一張圖表。" in text for text in fake.document_inputs())
    for names in fake.loaded_during_embed:
        assert NINE_B_MODEL not in names
        assert FOUR_B not in names
    assert fake.loaded == []

    store = open_store(config)
    index = ChromaIndex.open(config.data_dir)
    try:
        saved = store.list_documents()[0]
        assert saved.status == "complete"
        assert saved.parser_version == PARSER_VERSION
        chunks = store.chunks_for_document(saved.doc_id)
        figure = next(chunk for chunk in chunks if chunk.kind == "figure")
        assert figure.embed_state == "done"
        assert figure.page_start is None
        assert figure.page_end is None
        assert "描述：這是一張圖表。" in figure.text
        assert all(chunk.embed_state == "done" for chunk in chunks)
        image_row = store.images_for_document(saved.doc_id)[0]
        assert image_row.route == "label_and_caption"
        assert image_row.status == "caption_done"
        assert image_row.ocr_text == "圖表說明"
        vector = store.get_vector(figure.chunk_id, EMBED)
        assert vector is not None
        assert vector.native_dim == NATIVE
        assert len(vector.vector) == NATIVE
        active = store.get_active()
        assert active is not None
        assert active.collection == collection_name(EMBED, WIDTH)
        assert active.embed_model == EMBED
        assert active.dimensions == WIDTH
        found = index.get_embeddings(
            [figure.chunk_id],
            model=EMBED,
            dimensions=WIDTH,
            embed_num_ctx=2048,
        )
        assert len(found[figure.chunk_id]) == NATIVE
        assert index.collection_names() == [collection_name(EMBED, WIDTH)]
        finished = store.recent_runs()[0]
        assert finished.command == "ingest"
        assert finished.exit_code == EXIT_OK
        assert finished.ask_latency_ms is None
        assert finished.refused_count == 0
        assert finished.model_name == EMBED
        assert finished.ocr_count >= 1
        assert finished.vision_count >= 1
        assert finished.chunk_count >= 1
        assert "這是一張圖表" not in repr(finished)
        assert "前文在圖的前面" not in repr(finished)
    finally:
        index.close()
        store.close()

    log_path = config.data_dir / "logs" / "rag.log"
    log = log_path.read_text(encoding="utf-8")
    assert stat.S_IMODE(log_path.stat().st_mode) == 0o600
    for hidden in ("前文在圖的前面", "這是一張圖表", "圖表說明"):
        assert hidden not in log
        assert hidden not in captured.out
        assert hidden not in captured.err
    assert "階段 預檢" in log
    assert "階段 解析" in log
    assert "階段 OCR" in log
    assert "階段 圖像區塊" in log
    assert "階段 視覺" in log
    assert "階段 嵌入" in log
    assert "已入庫 note.docx" in captured.out


def json_body(request: httpx.Request):
    if not request.content:
        return None
    return json.loads(request.content.decode())


def test_text_rich_ocr_does_not_call_vision(tmp_path):
    config = write_rag(tmp_path)
    path = tmp_path / "shot.png"
    write_png(path)
    body = "額" * 40
    fake = IngestFake()
    code = run(config, fake, path, ocr_engine=ocr_engine(body))
    assert code == EXIT_OK
    assert "/api/chat" not in fake.paths()
    store = open_store(config)
    try:
        saved = store.list_documents()[0]
        chunks = store.chunks_for_document(saved.doc_id)
        assert len(chunks) == 1
        assert chunks[0].kind == "ocr"
        assert chunks[0].embed_state == "done"
        assert chunks[0].text.startswith("[圖像 OCR] shot.png")
        assert "第" not in chunks[0].text.splitlines()[0]
        assert body in chunks[0].text
        assert body in fake.document_inputs()[0]
    finally:
        store.close()


def test_caption_only_failure_deletes_the_placeholder(tmp_path):
    config = write_rag(tmp_path)
    png = tmp_path / "shot.png"
    write_png(png)
    document = tmp_path / "note.docx"
    write_docx(document, png)
    fake = IngestFake()
    fake.replies = [
        {"role": "assistant", "content": "", "thinking": "隱藏"},
        {"role": "assistant", "content": "   ", "thinking": ""},
    ]
    code = run(config, fake, document, ocr_engine=ocr_engine(""))
    assert code == EXIT_OK
    store = open_store(config)
    try:
        saved = store.list_documents()[0]
        chunks = store.chunks_for_document(saved.doc_id)
        assert chunks
        assert all(chunk.kind == "text" for chunk in chunks)
        assert all(chunk.embed_state == "done" for chunk in chunks)
        image_row = store.images_for_document(saved.doc_id)[0]
        assert image_row.status == "caption_failed"
        assert image_row.route == "caption_only"
    finally:
        store.close()


@pytest.mark.parametrize("strict, expected", [(False, EXIT_OK), (True, EXIT_NO_CHUNKS)])
def test_label_failure_keeps_the_chunk(tmp_path, capsys, strict, expected):
    config = write_rag(tmp_path)
    png = tmp_path / "shot.png"
    write_png(png)
    document = tmp_path / "note.docx"
    write_docx(document, png)
    fake = IngestFake()
    fake.replies = [
        {"role": "assistant", "content": "", "thinking": "隱藏"},
        {"role": "assistant", "content": "", "thinking": ""},
    ]
    code = run(config, fake, document, ocr_engine=ocr_engine("短標"), strict=strict)
    captured = capsys.readouterr()
    assert code == expected
    store = open_store(config)
    try:
        saved = store.list_documents()[0]
        chunks = store.chunks_for_document(saved.doc_id)
        figure = next(chunk for chunk in chunks if chunk.kind == "figure")
        text = next(chunk for chunk in chunks if "前文" in chunk.text)
        assert "描述" not in figure.text
        assert "圖中文字：短標" in figure.text
        assert figure.embed_state == "done"
        assert text.embed_state == "done"
        assert text.text in fake.document_inputs()
        assert figure.text in fake.document_inputs()
        assert store.images_for_document(saved.doc_id)[0].status == "caption_failed"
    finally:
        store.close()
    if strict:
        assert "有圖像或區塊失敗" in captured.err
        assert "已入庫" not in captured.out
    else:
        assert "已入庫 note.docx" in captured.out


def test_parser_version_change_reparses_and_drops_old_ids(tmp_path, monkeypatch):
    config = write_rag(tmp_path)
    path = tmp_path / "note.txt"
    path.write_text("原始正文。", encoding="utf-8")
    calls = count_parse(monkeypatch)
    fake = IngestFake()
    assert run(config, fake, path) == EXIT_OK
    store = open_store(config)
    index = ChromaIndex.open(config.data_dir)
    try:
        saved = store.list_documents()[0]
        original = store.chunks_for_document(saved.doc_id)[0]
        store.update_chunk_text(original.chunk_id, "OLD_BODY_SENTINEL", 4)
        extra = Chunk(
            chunk_id=chunk_id_for(saved.doc_id, original.ordinal + 1),
            doc_id=saved.doc_id,
            ordinal=original.ordinal + 1,
            text="EXTRA_CHROMA",
            kind="text",
            token_estimate=4,
            embed_state="done",
        )
        store.insert_chunk(extra)
        index.upsert(
            [
                IndexRecord(
                    chunk_id=extra.chunk_id,
                    doc_id=saved.doc_id,
                    text="EXTRA_CHROMA",
                    embedding=VECTOR,
                    source_path=saved.source_path,
                    filename=saved.filename,
                    kind="text",
                    ordinal=extra.ordinal,
                )
            ],
            model=EMBED,
            dimensions=WIDTH,
            embed_num_ctx=2048,
        )
        extra_id = extra.chunk_id
        doc_id = saved.doc_id
    finally:
        index.close()
        store.close()
    with sqlite3.connect(config.data_dir / "corpus.sqlite") as connection:
        connection.execute("UPDATE documents SET parser_version = '0'")
    assert run(config, fake, path) == EXIT_OK
    assert calls["n"] == 2
    store = open_store(config)
    index = ChromaIndex.open(config.data_dir)
    try:
        saved = store.get_document(doc_id)
        assert saved is not None
        assert saved.parser_version == PARSER_VERSION
        texts = [chunk.text for chunk in store.chunks_for_document(doc_id)]
        assert "OLD_BODY_SENTINEL" not in texts
        assert "EXTRA_CHROMA" not in texts
        assert any("原始正文" in text for text in texts)
        found = index.get_embeddings(
            [extra_id],
            model=EMBED,
            dimensions=WIDTH,
            embed_num_ctx=2048,
        )
        assert extra_id not in found
    finally:
        index.close()
        store.close()


def test_resume_does_not_embed_a_stale_missing_file(tmp_path):
    config = write_rag(tmp_path)
    store = open_store(config)
    try:
        seed_document(
            store,
            DOC_A,
            str(tmp_path / "gone.txt"),
            "STALE_SHOULD_NOT_EMBED",
            parser_version="0",
        )
        seed_document(store, DOC_B, str(tmp_path / "current.txt"), "CURRENT_SHOULD_EMBED")
    finally:
        store.close()
    fake = IngestFake()
    with pytest.raises(RagError) as exc:
        run(config, fake, None, resume=True, no_vision=True)
    assert exc.value.code == EXIT_NOT_FOUND
    assert "STALE_SHOULD_NOT_EMBED" not in fake.document_inputs()
    assert "CURRENT_SHOULD_EMBED" in fake.document_inputs()
    store = open_store(config)
    try:
        assert store.get_chunk(chunk_id_for(DOC_A, 0)).embed_state == "pending"
        assert store.get_chunk(chunk_id_for(DOC_B, 0)).embed_state == "done"
        assert store.get_document(DOC_A) is not None
    finally:
        store.close()


def test_resume_of_current_version_embeds_without_parsing(tmp_path, monkeypatch):
    config = write_rag(tmp_path)
    path = tmp_path / "note.txt"
    path.write_text("稍後再嵌入。", encoding="utf-8")
    calls = count_parse(monkeypatch)
    fake = IngestFake()
    assert run(config, fake, path, no_embed=True, no_vision=True) == EXIT_OK
    assert calls["n"] == 1
    assert fake.seen == []
    assert run(config, fake, None, resume=True, no_vision=True) == EXIT_OK
    assert calls["n"] == 1
    assert "稍後再嵌入。" in fake.document_inputs()
    store = open_store(config)
    try:
        chunk = store.chunks_for_document(store.list_documents()[0].doc_id)[0]
        assert chunk.embed_state == "done"
    finally:
        store.close()


def test_corrupt_docx_exits_5_without_a_document(tmp_path):
    config = write_rag(tmp_path)
    path = tmp_path / "bad.docx"
    path.write_bytes(b"not a docx")
    fake = IngestFake()
    with pytest.raises(RagError) as exc:
        run(config, fake, path, no_embed=True, no_vision=True)
    assert exc.value.code == EXIT_DOCUMENT
    assert fake.seen == []
    store = open_store(config)
    try:
        assert store.list_documents() == []
    finally:
        store.close()


def test_document_error_after_image_write_deletes_the_files(tmp_path, monkeypatch):
    config = write_rag(tmp_path)
    path = tmp_path / "shot.png"
    write_png(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()

    def explode(self, document, chunks, images):
        del self, document, chunks, images
        raise RagError(EXIT_DOCUMENT, "提交失敗。")

    monkeypatch.setattr(SqliteStore, "commit_parsed", explode)
    fake = IngestFake()
    with pytest.raises(RagError) as exc:
        run(config, fake, path)
    assert exc.value.code == EXIT_DOCUMENT
    assert not (config.data_dir / "images" / digest).exists()
    store = open_store(config)
    try:
        assert store.list_documents() == []
    finally:
        store.close()


def test_low_disk_exits_10_without_calling_ollama(tmp_path):
    config = write_rag(tmp_path, min_free=50)
    path = tmp_path / "note.txt"
    path.write_text("不會入庫。", encoding="utf-8")
    fake = IngestFake()
    with pytest.raises(RagError) as exc:
        run(config, fake, path, free_bytes=lambda: 1)
    assert exc.value.code == EXIT_DISK
    assert "1" in exc.value.message
    assert "50" in exc.value.message
    assert fake.seen == []
    assert not config.data_dir.exists()


def test_empty_text_exits_8_twice(tmp_path, capsys):
    config = write_rag(tmp_path)
    path = tmp_path / "blank.txt"
    path.write_text(" \n", encoding="utf-8")
    fake = IngestFake()
    assert run(config, fake, path, no_vision=True) == EXIT_NO_CHUNKS
    first = capsys.readouterr()
    assert "沒有可嵌入的區塊" in first.err
    assert "已經入庫" not in first.out
    assert run(config, fake, path, no_vision=True) == EXIT_NO_CHUNKS
    second = capsys.readouterr()
    assert "沒有可嵌入的區塊" in second.err
    assert "已經入庫" not in second.out
    store = open_store(config)
    try:
        documents = store.list_documents()
        assert len(documents) == 1
        assert documents[0].status == "complete"
        assert store.chunks_for_document(documents[0].doc_id) == []
    finally:
        store.close()


def test_png_without_ocr_or_vision_exits_8(tmp_path, capsys):
    config = write_rag(tmp_path)
    path = tmp_path / "shot.png"
    write_png(path)
    fake = IngestFake()
    code = run(config, fake, path, no_ocr=True, no_vision=True)
    captured = capsys.readouterr()
    assert code == EXIT_NO_CHUNKS
    assert "有圖像或區塊失敗" in captured.err
    assert "/api/chat" not in fake.paths()
    store = open_store(config)
    try:
        saved = store.list_documents()[0]
        assert store.chunks_for_document(saved.doc_id) == []
        assert store.images_for_document(saved.doc_id)[0].status == "empty_no_vision"
    finally:
        store.close()


def test_no_ocr_rejects_more_than_20_images_before_commit(tmp_path, monkeypatch):
    config = write_rag(tmp_path)
    path = tmp_path / "note.txt"
    path.write_text("x", encoding="utf-8")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    resolved = str(path.resolve())

    def crowd(_path, _limits=None):
        images = []
        for ordinal in range(21):
            image_id = image_id_for(digest, ordinal)
            images.append(
                ParsedImage(
                    image_id=image_id,
                    doc_id=digest,
                    storage_relpath=storage_relpath(digest, image_id),
                    width=80,
                    height=80,
                    byte_size=16,
                    sha256=hashlib.sha256(f"img{ordinal}".encode()).hexdigest(),
                    kind="embedded",
                    status="pending_ocr",
                    preceding_chunk_count=0,
                )
            )
        return ParseResult(
            ParsedDocument(
                doc_id=digest,
                source_path=resolved,
                filename=path.name,
                media_type="text",
                byte_size=path.stat().st_size,
                sha256=digest,
                parser="text",
                parser_version=PARSER_VERSION,
            ),
            (),
            tuple(images),
        )

    monkeypatch.setattr("ragcli.parse.parse_file", crowd)
    fake = IngestFake()
    with pytest.raises(RagError) as exc:
        run(config, fake, path, no_ocr=True, no_embed=True, no_vision=True)
    assert exc.value.code == EXIT_USAGE
    assert exc.value.message == "--no-ocr 時圖像超過 20 張。"
    assert fake.seen == []
    store = open_store(config)
    try:
        assert store.list_documents() == []
    finally:
        store.close()


def test_resume_no_ocr_leaves_a_stored_crowd(tmp_path):
    config = write_rag(tmp_path)
    store = open_store(config)
    try:
        seed_document(store, DOC_C, str(tmp_path / "crowd.txt"), "保留的文字。", embed_state="done")
        for ordinal in range(21):
            image_id = image_id_for(DOC_C, ordinal)
            store.insert_image(
                ImageRecord(
                    image_id=image_id,
                    doc_id=DOC_C,
                    storage_relpath=storage_relpath(DOC_C, image_id),
                    width=80,
                    height=80,
                    byte_size=16,
                    sha256=hashlib.sha256(f"crowd{ordinal}".encode()).hexdigest(),
                    kind="embedded",
                    status="pending_ocr",
                )
            )
    finally:
        store.close()
    fake = IngestFake()
    with pytest.raises(RagError) as exc:
        run(config, fake, None, resume=True, no_ocr=True, no_vision=True)
    assert exc.value.code == EXIT_USAGE
    assert fake.seen == []
    store = open_store(config)
    try:
        assert len(store.images_for_document(DOC_C)) == 21
        assert store.chunks_for_document(DOC_C)[0].text == "保留的文字。"
    finally:
        store.close()


def test_missing_embed_model_exits_7(tmp_path):
    config = write_rag(tmp_path)
    path = tmp_path / "note.txt"
    path.write_text("需要嵌入模型。", encoding="utf-8")
    fake = IngestFake()
    fake.tag_names = [NINE_B_MODEL, FOUR_B]
    with pytest.raises(RagError) as exc:
        run(config, fake, path, no_vision=True)
    assert exc.value.code == EXIT_MODEL_MISSING
    assert EMBED in exc.value.message
    assert "rag models pull" in exc.value.message
    assert "ollama pull" not in exc.value.message
    index = ChromaIndex.open(config.data_dir)
    try:
        assert index.collection_names() == []
    finally:
        index.close()


def test_active_model_missing_does_not_create_the_configured_collection(tmp_path):
    config = write_rag(tmp_path)
    path = tmp_path / "note.txt"
    path.write_text("作用中模型。", encoding="utf-8")
    store = open_store(config)
    try:
        store.set_active(
            ActiveIndex(
                collection=collection_name("qwen3-embedding:4b", 2560),
                embed_model="qwen3-embedding:4b",
                dimensions=2560,
                embed_num_ctx=2048,
            )
        )
    finally:
        store.close()
    fake = IngestFake()
    with pytest.raises(RagError) as exc:
        run(config, fake, path, no_vision=True)
    assert exc.value.code == EXIT_MODEL_MISSING
    assert "qwen3-embedding:4b" in exc.value.message
    assert EMBED not in exc.value.message
    assert "rag models pull" in exc.value.message
    assert "ollama pull" not in exc.value.message
    index = ChromaIndex.open(config.data_dir)
    try:
        assert index.collection_names() == []
    finally:
        index.close()


def test_active_shorter_width_truncates_without_changing_the_pointer(tmp_path):
    config = write_rag(tmp_path)
    path = tmp_path / "note.txt"
    path.write_text("截斷寬度。", encoding="utf-8")
    narrow = collection_name(EMBED, 2)
    store = open_store(config)
    try:
        store.set_active(
            ActiveIndex(
                collection=narrow,
                embed_model=EMBED,
                dimensions=2,
                embed_num_ctx=2048,
            )
        )
    finally:
        store.close()
    fake = IngestFake()
    assert run(config, fake, path, no_vision=True) == EXIT_OK
    store = open_store(config)
    index = ChromaIndex.open(config.data_dir)
    try:
        saved = store.list_documents()[0]
        chunk = store.chunks_for_document(saved.doc_id)[0]
        vector = store.get_vector(chunk.chunk_id, EMBED)
        assert vector is not None
        assert len(vector.vector) == NATIVE
        assert vector.vector[0] == pytest.approx(1.0)
        assert vector.vector[3] == pytest.approx(4.0)
        found = index.get_embeddings(
            [chunk.chunk_id],
            model=EMBED,
            dimensions=2,
            embed_num_ctx=2048,
        )
        assert len(found[chunk.chunk_id]) == 2
        assert found[chunk.chunk_id][0] == pytest.approx(1.0)
        assert found[chunk.chunk_id][1] == pytest.approx(0.0)
        active = store.get_active()
        assert active is not None
        assert active.collection == narrow
        assert active.dimensions == 2
        assert active.embed_model == EMBED
        assert collection_name(EMBED, NATIVE) not in index.collection_names()
        assert index.collection_names() == [narrow]
    finally:
        index.close()
        store.close()


def test_embed_batches_skip_blocked_chunks(tmp_path):
    config = write_rag(tmp_path, batch_size=2)
    store = open_store(config)
    try:
        store.insert_document(
            Document(
                doc_id=DOC_C,
                source_path=str(tmp_path / "batches.txt"),
                filename="batches.txt",
                media_type="text",
                byte_size=8,
                sha256=DOC_C,
                parser="text",
                parser_version=PARSER_VERSION,
                ingested_at=WHEN,
                status="parsed",
            )
        )
        rows = (
            (0, "pending", "第一批甲"),
            (1, "pending", "第一批乙"),
            (2, "pending", "第二批丙"),
            (3, "blocked", "BLOCKED_SECRET"),
        )
        for ordinal, state, text in rows:
            store.insert_chunk(
                Chunk(
                    chunk_id=chunk_id_for(DOC_C, ordinal),
                    doc_id=DOC_C,
                    ordinal=ordinal,
                    text=text,
                    kind="text",
                    token_estimate=4,
                    embed_state=state,
                )
            )
    finally:
        store.close()
    fake = IngestFake()
    assert run(config, fake, None, resume=True, no_vision=True) == EXIT_OK
    assert [len(batch) for batch in fake.embed_inputs] == [2, 1]
    assert fake.embed_inputs[0] == ["第一批甲", "第一批乙"]
    assert fake.embed_inputs[1] == ["第二批丙"]
    assert "BLOCKED_SECRET" not in fake.document_inputs()
    body = fake.embed_bodies()[0]
    assert body["truncate"] is False
    assert "dimensions" not in body
    assert body["options"]["num_ctx"] == 2048
    assert body["keep_alive"] == "30m"
    assert not body["input"][0].startswith("Instruct:")
    store = open_store(config)
    try:
        chunks = store.chunks_for_document(DOC_C)
        assert [chunk.embed_state for chunk in chunks] == ["done", "done", "done", "blocked"]
    finally:
        store.close()


def test_same_hash_updates_a_moved_path_without_reparsing(tmp_path, monkeypatch, capsys):
    config = write_rag(tmp_path)
    path = tmp_path / "note.txt"
    path.write_text("路徑會改。", encoding="utf-8")
    calls = count_parse(monkeypatch)
    fake = IngestFake()
    assert run(config, fake, path, no_vision=True) == EXIT_OK
    assert "已入庫 note.txt" in capsys.readouterr().out
    assert run(config, fake, path, no_vision=True) == EXIT_OK
    assert "已經入庫 note.txt" in capsys.readouterr().out
    moved = tmp_path / "moved-note.txt"
    path.rename(moved)
    assert run(config, fake, moved, no_vision=True) == EXIT_OK
    assert calls["n"] == 1
    assert "已經入庫 moved-note.txt" in capsys.readouterr().out
    store = open_store(config)
    index = ChromaIndex.open(config.data_dir)
    try:
        saved = store.list_documents()[0]
        assert saved.source_path == str(moved.resolve())
        assert saved.filename == "moved-note.txt"
        hits = index.query(
            VECTOR,
            model=EMBED,
            dimensions=WIDTH,
            embed_num_ctx=2048,
            n_results=4,
        )
        assert hits
        assert all(hit.source_path == str(moved.resolve()) for hit in hits)
        assert all(hit.filename == "moved-note.txt" for hit in hits)
    finally:
        index.close()
        store.close()


def test_force_reparses(tmp_path, monkeypatch):
    config = write_rag(tmp_path)
    path = tmp_path / "note.txt"
    path.write_text("強制重來。", encoding="utf-8")
    calls = count_parse(monkeypatch)
    fake = IngestFake()
    assert run(config, fake, path, no_vision=True) == EXIT_OK
    assert run(config, fake, path, no_vision=True, force=True) == EXIT_OK
    assert calls["n"] == 2
    store = open_store(config)
    try:
        assert len(store.list_documents()) == 1
    finally:
        store.close()


def test_no_model_flags_do_not_call_ollama_or_log_the_body(tmp_path, capsys):
    config = write_rag(tmp_path)
    sentence = "僅供日誌檢查的句子。"
    path = tmp_path / "note.txt"
    path.write_text(sentence, encoding="utf-8")
    fake = IngestFake()
    logger = logging.getLogger("ragcli")
    previous = logger.level
    logger.setLevel(logging.INFO)
    try:
        code = run(config, fake, path, no_embed=True, no_vision=True)
    finally:
        logger.setLevel(previous)
    captured = capsys.readouterr()
    assert code == EXIT_OK
    assert fake.seen == []
    log_path = config.data_dir / "logs" / "rag.log"
    log = log_path.read_text(encoding="utf-8")
    assert sentence not in log
    assert sentence not in captured.out
    assert sentence not in captured.err
    assert "階段 解析" in log
    assert stat.S_IMODE(log_path.stat().st_mode) == 0o600
    assert "已入庫 note.txt" in captured.out


def test_skipped_bomb_does_not_call_ocr(tmp_path, monkeypatch):
    config = write_rag(tmp_path)
    png = tmp_path / "shot.png"
    write_png(png, 30, 20)
    document = tmp_path / "bomb.docx"
    write_docx_line = DocxDocument()
    write_docx_line.add_paragraph("圖旁邊仍有文字。")
    write_docx_line.add_picture(str(png))
    write_docx_line.save(document)
    replace_media(document, bomb_png(8000, 5000))
    monkeypatch.setattr(
        "ragcli.parse.pixels.Image.open",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("decoded")),
    )
    fake = IngestFake()
    assert run(config, fake, document) == EXIT_OK
    assert "/api/chat" not in fake.paths()
    store = open_store(config)
    try:
        saved = store.list_documents()[0]
        chunks = store.chunks_for_document(saved.doc_id)
        assert any("圖旁邊仍有文字" in chunk.text for chunk in chunks)
        assert all(chunk.kind != "ocr" for chunk in chunks)
        image_row = store.images_for_document(saved.doc_id)[0]
        assert image_row.status == "skipped_bomb"
        assert image_row.route == "skipped_bomb"
    finally:
        store.close()


def test_ollama_down_exits_6_without_a_document(tmp_path):
    config = write_rag(tmp_path)
    path = tmp_path / "note.txt"
    path.write_text("連不上。", encoding="utf-8")
    fake = IngestFake()
    fake.down = True
    with pytest.raises(RagError) as exc:
        run(config, fake, path)
    assert exc.value.code == EXIT_OLLAMA_UNREACHABLE
    assert "連不上 Ollama" in exc.value.message
    store = open_store(config)
    try:
        assert store.list_documents() == []
        finished = store.recent_runs()[0]
        assert finished.exit_code == EXIT_OLLAMA_UNREACHABLE
        assert finished.ollama_error_count >= 1
        assert finished.command == "ingest"
    finally:
        store.close()


def test_stuck_generator_exits_13_before_embed(tmp_path):
    config = write_rag(tmp_path)
    path = tmp_path / "note.txt"
    path.write_text("生成模型佔著。", encoding="utf-8")
    fake = IngestFake()
    fake.stuck = True
    fake.loaded = [{"name": NINE_B_MODEL, "context_length": 4096}]
    with pytest.raises(RagError) as exc:
        run(config, fake, path, no_vision=True)
    assert exc.value.code == EXIT_RESOURCES
    assert "卸載後仍有生成模型" in exc.value.message
    assert fake.document_inputs() == []

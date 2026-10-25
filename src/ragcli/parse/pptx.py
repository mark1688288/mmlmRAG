"""PPTX parser.

Shapes are read in order. Picture blobs go through the header pixel check
before they are decoded. Notes become their own text with section ``notes``.
The parser does not write files.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ragcli.errors import EXIT_DOCUMENT, RagError
from ragcli.parse.docx import _classify
from ragcli.parse.model import ParseLimits, ParseResult, Segment, TextPiece, finish, load_bytes


def parse_pptx(path: Path, limits: ParseLimits | None = None) -> ParseResult:
    """Parse one PPTX into a document, text chunks, and image records."""
    limits = limits or ParseLimits()
    resolved, data = load_bytes(path, limits.max_bytes)
    try:
        presentation = Presentation(str(resolved))
    except Exception as exc:
        raise RagError(EXIT_DOCUMENT, "檔案已損壞，無法解析。") from exc
    if len(presentation.slides) > limits.max_pages:
        raise RagError(EXIT_DOCUMENT, "頁數超過上限。")
    segments: list[Segment] = []
    for index, slide in enumerate(presentation.slides, start=1):
        _read_slide(slide, index, segments, limits)
    return finish(
        resolved=resolved,
        data=data,
        media_type="pptx",
        parser="python-pptx",
        title=None,
        page_count=len(presentation.slides),
        segments=segments,
        limits=limits,
    )


def _read_slide(slide, page: int, segments: list[Segment], limits: ParseLimits) -> None:
    _shapes(slide.shapes, page, segments, limits, section=None)
    notes = _notes(slide)
    if notes:
        segments.append(TextPiece(notes, page, "notes"))


def _shapes(shapes, page: int, segments: list[Segment], limits: ParseLimits, section: str | None) -> None:
    buffer: list[str] = []

    def flush() -> None:
        text = "\n".join(part for part in buffer if part.strip()).strip()
        buffer.clear()
        if text:
            segments.append(TextPiece(text, page, section))

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            flush()
            _shapes(shape.shapes, page, segments, limits, section)
            continue
        if getattr(shape, "has_table", False):
            text = _table_text(shape.table)
            if text:
                buffer.append(text)
        elif getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            if text:
                buffer.append(text)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            flush()
            try:
                blob = shape.image.blob
            except Exception as exc:
                raise RagError(EXIT_DOCUMENT, "檔案已損壞，無法解析。") from exc
            segments.append(_classify(blob, page=page, limits=limits, locator=f"pptx:{page}"))
    flush()


def _table_text(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.replace("\t", " ").replace("\n", " ").strip() for cell in row.cells]
        rows.append("\t".join(cells))
    return "\n".join(rows).strip()


def _notes(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    try:
        text = slide.notes_slide.notes_text_frame.text
    except Exception:
        return ""
    return text.strip()
